"""Lectura de documentos y paquetes del expediente; nunca ejecuta contenido subido."""
from __future__ import annotations

import csv
import io
import json
import re
import uuid
import zipfile
from hashlib import sha256

from . import workspace_store as store
from .workspace import normalize, prepare

MAX_BYTES = 10 * 1024 * 1024
MAX_TEXT = 500000


def table(content, name):
    if name.lower().endswith(".xlsx"):
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 60 * 1024 * 1024:
                raise ValueError("La hoja descomprimida excede el límite de lectura")
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            return [["" if v is None else str(v) for v in row] for row in workbook.active.iter_rows(values_only=True)]
        finally:
            workbook.close()
    text = content.decode("utf-8-sig")
    if "\0" in text:
        raise ValueError("El archivo no contiene texto UTF-8 legible")
    try:
        dialect = csv.Sniffer().sniff(text[:8000], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return list(csv.reader(io.StringIO(text), dialect))


def extract(content, name):
    suffix = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    pages, text = [], ""
    if suffix in ("txt", "md", "csv"):
        text = content.decode("utf-8-sig")
        if "\0" in text:
            raise ValueError("No es texto legible")
    elif suffix == "xlsx":
        rows = table(content, name)
        text = "\n".join(f"Fila {i}: " + " | ".join(row) for i, row in enumerate(rows, 1))
    elif suffix == "pdf":
        if not content.startswith(b"%PDF-"):
            raise ValueError("El archivo no tiene una cabecera PDF válida")
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content))
        for i, page in enumerate(reader.pages):
            part = page.extract_text() or ""
            pages.append({"page": i + 1, "offset": len(text), "length": len(part)})
            text += part + "\n"
            if len(text) > MAX_TEXT:
                break
    return {"text": text[:MAX_TEXT], "truncated": len(text) > MAX_TEXT, "pages": pages,
        "status": "text_extracted" if text.strip() else "pending"}


def upload(workspace_id, name, topic, content, content_type=""):
    if len(content) > MAX_BYTES:
        raise ValueError("El documento supera los 10 MB")
    name = re.split(r"[/\\]", name)[-1][:220]
    if not name:
        raise ValueError("El documento necesita nombre")
    try:
        extracted = extract(content, name)
    except Exception as exc:
        extracted = {"status": "pending", "error": f"No se pudo extraer texto: {str(exc)[:180]}"}
    metadata = {"id": uuid.uuid4().hex, "topic": topic, "name": name, "type": content_type,
        "size": len(content), "date": store.now(), "origin": "usuario", "storage": "server", "sha256": sha256(content).hexdigest(), **extracted}
    return store.put_document(workspace_id, metadata, content)


OPERATOR_HEADERS = {
    "id": ("id", "id operador", "codigo", "numero empleado"), "name": ("nombre", "operador", "nombre de operador"),
    "station": ("estacion", "area", "estacion o area"), "machine": ("maquina", "equipo", "maquina que usa"),
    "shift": ("turno",), "efficiency_reported": ("eficiencia", "eficiencia reportada"),
    "pace_rating": ("ritmo", "valoracion de ritmo", "valoracion"), "period": ("periodo", "fecha"),
}


def operator_preview(workspace_id, source_id, mapping=None):
    metadata, content = store.document(workspace_id, source_id)
    if not metadata["name"].lower().endswith((".csv", ".xlsx")):
        raise ValueError("La base de operadores debe ser CSV o XLSX")
    rows = table(content, metadata["name"])
    if not rows:
        raise ValueError("La tabla está vacía")
    headers = rows[0]
    aliases = {normalize(h): i for i, h in enumerate(headers)}
    columns = {}
    for key, names in OPERATOR_HEADERS.items():
        requested = (mapping or {}).get(key)
        index = aliases.get(normalize(requested)) if requested else next((aliases[normalize(n)] for n in names if normalize(n) in aliases), None)
        columns[key] = index
    if columns["name"] is None:
        raise ValueError("Indica la columna del nombre de operador")
    result, errors = [], []
    for line, row in enumerate(rows[1:], 2):
        if not any(str(v).strip() for v in row):
            continue
        item = {key: str(row[index]).strip() if index is not None and index < len(row) else "" for key, index in columns.items()}
        if not item["name"]:
            errors.append({"row": line, "error": "Falta nombre de operador"})
            continue
        item["operator_code"] = item.pop("id")
        item.update(id=f"operador-{source_id}-{line}", source_id=source_id, row=line, original=row,
            status="pending", stationId="", machineId="", calculated_efficiency=None)
        for metric in ("efficiency_reported", "pace_rating"):
            raw = item[metric]
            try:
                value = float(raw.strip("% ").replace(",", ".")) if raw else None
                if value is not None and (value < 0 or not value < float("inf")):
                    raise ValueError()
            except ValueError:
                value = None
                errors.append({"row": line, "error": f"Revisar {metric}: {raw}"})
            item[metric] = {"raw": raw, "value": value, "scale": "percent" if "%" in raw else "unspecified", "source": "imported"}
        result.append(item)
    return {"headers": headers, "mapping": {k: headers[v] if v is not None else None for k, v in columns.items()}, "operators": result, "errors": errors,
        "note": "Indicadores importados. La eficiencia calculada necesita observaciones y una definición; no modifica el factor global del motor."}


def export_bundle(workspace_id):
    saved = store.read(workspace_id)
    result = io.BytesIO()
    referenced = {f["id"] for f in saved["project"].get("contextWorkspace", {}).get("attachments", [])}
    docs = [d for d in store.documents(workspace_id) if d["id"] in referenced]
    if len(docs) > 998 or sum(d["size"] for d in docs) > 90 * 1024 * 1024:
        raise ValueError("El paquete excede 90 MB de originales o 998 archivos. Divide el proyecto antes de exportar.")
    with zipfile.ZipFile(result, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("manifest.json", store.encode({"format": "cauce-workspace-1", "project": saved["project"], "documents": docs}))
        for doc in docs:
            archive.writestr(f"documents/{doc['id']}", store.document(workspace_id, doc["id"])[1])
    return result.getvalue()


def import_bundle(content):
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        infos = archive.infolist()
        if len(infos) > 1000 or sum(i.file_size for i in infos) > 100 * 1024 * 1024:
            raise ValueError("El paquete supera el límite de 100 MB descomprimidos o 1000 archivos")
        manifest = json.loads(archive.read("manifest.json"))
        if manifest.get("format") != "cauce-workspace-1" or not isinstance(manifest.get("project"), dict):
            raise ValueError("Paquete de proyecto no reconocido")
        workspace_id = uuid.uuid4().hex
        project = prepare(manifest["project"])
        mappings, validated = {}, []
        for metadata in manifest.get("documents", []):
            original = metadata["id"]
            if not re.fullmatch(r"[\w-]+", original, re.ASCII):
                raise ValueError("Referencia de archivo inválida")
            binary = archive.read(f"documents/{original}")
            if len(binary) > MAX_BYTES or sha256(binary).hexdigest() != metadata["sha256"]:
                raise ValueError("Archivo demasiado grande o contenido alterado")
            # IDs globales nuevos: restaurar una copia no pisa documentos anteriores.
            replacement = uuid.uuid4().hex
            mappings[original] = replacement
            validated.append(({**metadata, "id": replacement}, binary))
        def remap(value):
            if isinstance(value, dict):
                return {key: remap(item) for key, item in value.items()}
            if isinstance(value, list):
                return [remap(item) for item in value]
            return mappings.get(value, value) if isinstance(value, str) else value
        project = remap(project)
        saved = store.write(workspace_id, project, 0)
        for metadata, binary in validated:
            store.put_document(workspace_id, metadata, binary)
        return saved
